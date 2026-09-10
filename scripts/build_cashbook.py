# -*- coding: utf-8 -*-
"""現金出納帳ブックを「202606〜202705の1年分テンプレート」に組み直す。

・入金（収入金額）は扱わない
・月別シート12枚（日付／科目／摘要／支払金額）
・勘定科目シートで月ごと・科目ごとに自動集計
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SRC, DST = sys.argv[1], sys.argv[2]

MONTHS = ["202606", "202607", "202608", "202609", "202610", "202611",
          "202612", "202701", "202702", "202703", "202704", "202705"]

ACCOUNTS = ["雑費", "リース料", "賃借料", "接待交際費", "会議費", "通信費",
            "旅費交通費", "消耗品費", "租税公課", "修繕費", "水道光熱費",
            "福利厚生費", "支払手数料", "保険料", "広告宣伝費", "車輛燃料費",
            "工具、器具及び備品"]

ROWS_PER_SHEET = 200          # 入力用の空行を確保する行数
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
NUM = "#,##0_);[Red]\\(#,##0\\)"


def read_existing(wb, name):
    """既存シートから支払行だけを取り出す（入金行は捨てる）。"""
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    out = []
    for r in range(2, ws.max_row + 1):
        date, acc, memo = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        pay = ws.cell(r, 5).value
        if not isinstance(pay, (int, float)):     # 入金のみ／空行は除外
            continue
        out.append((date, ws.cell(r, 1).number_format,
                    str(acc).strip() if acc else "",
                    str(memo).strip() if memo else "", pay))
    return out


def main():
    src = openpyxl.load_workbook(SRC)
    data = {m: read_existing(src, m) for m in MONTHS}
    for rows in data.values():
        for _, _, acc, _, _ in rows:
            if acc and acc not in ACCOUNTS:
                ACCOUNTS.append(acc)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    last_acc_row = 1 + len(ACCOUNTS)
    dv_src = f"=勘定科目!$A$2:$A${last_acc_row}"

    for m in MONTHS:
        ws = wb.create_sheet(m)
        for i, h in enumerate(["日付", "科目", "摘要", "支払金額"], 1):
            c = ws.cell(1, i, h)
            c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
            c.alignment = Alignment(horizontal="center")
        for i, (date, fmt, acc, memo, pay) in enumerate(data[m]):
            r = i + 2
            c = ws.cell(r, 1, date)
            if hasattr(date, "year"):
                c.number_format = "yyyy/m/d"
            elif fmt:
                c.number_format = fmt
            ws.cell(r, 2, acc)
            ws.cell(r, 3, memo)
            c = ws.cell(r, 4, pay)
            c.number_format = NUM
        for r in range(2, ROWS_PER_SHEET + 2):
            ws.cell(r, 4).number_format = NUM
            for col in range(1, 5):
                ws.cell(r, col).border = BORDER
        dv = DataValidation(type="list", formula1=dv_src, allow_blank=True,
                            showDropDown=False)
        dv.error = "勘定科目シートに登録されている科目から選んでください"
        dv.errorTitle = "科目が一致しません"
        ws.add_data_validation(dv)
        dv.add(f"B2:B{ROWS_PER_SHEET + 1}")
        for col, w in [("A", 13), ("B", 16), ("C", 36), ("D", 14)]:
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:D{ROWS_PER_SHEET + 1}"

    # ---------------- 勘定科目（自動集計） ----------------
    sm = wb.create_sheet("勘定科目", 0)
    c = sm.cell(1, 1, "勘定科目")
    c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
    for j, m in enumerate(MONTHS):
        c = sm.cell(1, 2 + j, m)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(horizontal="center")
    c = sm.cell(1, 2 + len(MONTHS), "年間合計")
    c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER

    for k, acc in enumerate(ACCOUNTS):
        r = k + 2
        sm.cell(r, 1, acc).border = BORDER
        for j, m in enumerate(MONTHS):
            c = sm.cell(r, 2 + j, f"=SUMIF('{m}'!$B:$B,$A{r},'{m}'!$D:$D)")
            c.number_format, c.border = NUM, BORDER
        L2 = get_column_letter(1 + len(MONTHS))
        c = sm.cell(r, 2 + len(MONTHS), f"=SUM(B{r}:{L2}{r})")
        c.number_format, c.border, c.font = NUM, BORDER, Font(bold=True)

    tr = last_acc_row + 1
    sm.cell(tr, 1, "科目合計")
    for j in range(len(MONTHS) + 1):
        L = get_column_letter(2 + j)
        c = sm.cell(tr, 2 + j, f"=SUM({L}2:{L}{tr - 1})")
        c.number_format = NUM
    for col in range(1, 3 + len(MONTHS)):
        c = sm.cell(tr, col)
        c.fill, c.font, c.border = TOTAL_FILL, Font(bold=True), BORDER

    # 支払総額と、科目未入力などで拾えなかった差額のチェック行
    pr = tr + 1
    sm.cell(pr, 1, "支払総額（各月シートD列）").font = Font(bold=True)
    for j, m in enumerate(MONTHS):
        c = sm.cell(pr, 2 + j, f"=SUM('{m}'!$D:$D)")
        c.number_format, c.border = NUM, BORDER
    L2 = get_column_letter(1 + len(MONTHS))
    c = sm.cell(pr, 2 + len(MONTHS), f"=SUM(B{pr}:{L2}{pr})")
    c.number_format, c.border = NUM, BORDER
    dr = pr + 1
    sm.cell(dr, 1, "差額（科目未入力チェック）").font = Font(bold=True, color="C00000")
    for j in range(len(MONTHS) + 1):
        L = get_column_letter(2 + j)
        c = sm.cell(dr, 2 + j, f"={L}{pr}-{L}{tr}")
        c.number_format, c.border = NUM, BORDER

    note = dr + 2
    for i, t in enumerate([
        "※ 各月シート（202606〜202705）のB列「科目」とD列「支払金額」から自動集計しています。",
        "※ 行を追加するときは、各月シートの表の中（2行目以降）に入力してください。列全体を参照しているため何行でも集計されます。",
        "※ B列の科目はこのシートのA列から選択できます。科目を増やすときはA列に追記し、上の合計式の範囲を広げてください。",
        "※「差額」が0以外のときは、科目が未入力または表記ゆれの行があります。",
        "※ 入金（収入）は集計対象外です。",
    ]):
        sm.cell(note + i, 1, t).font = Font(size=9, color="808080")

    sm.column_dimensions["A"].width = 26
    for j in range(len(MONTHS) + 1):
        sm.column_dimensions[get_column_letter(2 + j)].width = 12
    sm.freeze_panes = "B2"
    wb.active = 0
    wb.save(DST)
    print({m: len(v) for m, v in data.items() if v})
    print("accounts:", len(ACCOUNTS), "sheets:", wb.sheetnames)


main()
