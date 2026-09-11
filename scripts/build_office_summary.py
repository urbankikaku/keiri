# -*- coding: utf-8 -*-
"""事務所別の支払ブックに、1枚で月ごとの支払額と合計が分かる「月次一覧」シートを追加する。

行＝支払月（各事務所シートと同じ並び）、列＝事務所。
振込額／請求額（税込み）／源泉徴収額 の3ブロックに、事務所ごとの金額と合計を並べ、
暦年ごとの年計と総合計を入れる。数値は各事務所シートへの参照なので、
元シートを直せば一覧も更新される。
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC, DST = sys.argv[1], sys.argv[2]

OFFICES = ["井坂事務所", "岡崎事務所", "綾測"]
# 表示するブロック（見出し, 元シートの列）
BLOCKS = [("振込額（支払額）", "F"),
          ("請求額（税込み）", "D"),
          ("源泉徴収額", "E")]
FIRST_SRC_ROW = 2
SHEET = "月次一覧"

THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="1F4E79")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")
SUM_FILL = PatternFill("solid", fgColor="DDEBF7")
YEAR_FILL = PatternFill("solid", fgColor="FFF2CC")
GRAND_FILL = PatternFill("solid", fgColor="F8CBAD")
NUM = "#,##0_);[Red]\\(#,##0\\)"


def main():
    wb = openpyxl.load_workbook(SRC)
    base = wb[OFFICES[0]]

    # 支払月の一覧（元シートの日付行をそのまま使う）
    months = []
    r = FIRST_SRC_ROW
    while base.cell(r, 1).value is not None:
        months.append((r, base.cell(r, 1).value))
        r += 1
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)

    n_off = len(OFFICES)
    width = n_off + 1                      # 事務所＋合計
    ws.cell(1, 1, "事務所別 月次支払一覧").font = Font(bold=True, size=14)
    ws.cell(1, 3, "※ 金額は各事務所シートへの参照です。元シートを直すとこの表も変わります。").font = \
        Font(size=9, color="808080")

    HR1, HR2, FIRST_ROW = 3, 4, 5
    ws.merge_cells(start_row=HR1, start_column=1, end_row=HR2, end_column=1)
    c = ws.cell(HR1, 1, "支払日")
    c.fill, c.font = HDR_FILL, HDR_FONT
    c.alignment = Alignment(horizontal="center", vertical="center")
    for bi, (title, _) in enumerate(BLOCKS):
        c0 = 2 + bi * width
        ws.merge_cells(start_row=HR1, start_column=c0, end_row=HR1, end_column=c0 + width - 1)
        c = ws.cell(HR1, c0, title)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        for j, name in enumerate(OFFICES + ["合計"]):
            c = ws.cell(HR2, c0 + j, name)
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
    for cc in range(1, 2 + len(BLOCKS) * width):
        ws.cell(HR1, cc).border = BORDER

    def block_cols(bi):
        c0 = 2 + bi * width
        return c0, c0 + n_off - 1, c0 + n_off      # 先頭, 事務所末尾, 合計列

    # ---- 月行（暦年ごとに年計を挟む） ----
    row = FIRST_ROW
    year_rows, month_rows = {}, []
    for idx, (src_row, dt) in enumerate(months):
        year = dt.year
        if idx and year != months[idx - 1][1].year:
            year_rows[months[idx - 1][1].year] = row
            row += 1
        c = ws.cell(row, 1, f"='{OFFICES[0]}'!A{src_row}")
        c.number_format = 'yyyy"年"m"月"'
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
        for bi, (_, col) in enumerate(BLOCKS):
            c0, cend, ctot = block_cols(bi)
            for oi, off in enumerate(OFFICES):
                ref = f"'{off}'!{col}{src_row}"
                c = ws.cell(row, c0 + oi, f'=IF(N({ref})=0,"",{ref})')
                c.number_format, c.border = NUM, BORDER
            L1, L2 = get_column_letter(c0), get_column_letter(cend)
            c = ws.cell(row, ctot, f'=IF(SUM({L1}{row}:{L2}{row})=0,"",SUM({L1}{row}:{L2}{row}))')
            c.number_format, c.border, c.font = NUM, BORDER, Font(bold=True)
            c.fill = SUM_FILL
        month_rows.append((year, row))
        row += 1
    year_rows[months[-1][1].year] = row
    row += 1
    grand_row = row

    # ---- 年計行 ----
    for year, yrow in year_rows.items():
        rows = [r for y, r in month_rows if y == year]
        c = ws.cell(yrow, 1, f"{year}年 計")
        for cc in range(2, 2 + len(BLOCKS) * width):
            L = get_column_letter(cc)
            f = "+".join(f"{L}{r}" for r in rows)
            c2 = ws.cell(yrow, cc, f'=IF(SUM({L}{rows[0]}:{L}{rows[-1]})=0,"",'
                                   f'SUM({L}{rows[0]}:{L}{rows[-1]}))')
            c2.number_format = NUM
        for cc in range(1, 2 + len(BLOCKS) * width):
            cel = ws.cell(yrow, cc)
            cel.fill, cel.font, cel.border = YEAR_FILL, Font(bold=True), BORDER

    # ---- 総合計行 ----
    ws.cell(grand_row, 1, "総合計")
    for cc in range(2, 2 + len(BLOCKS) * width):
        L = get_column_letter(cc)
        f = "+".join(f"{L}{r}" for r in year_rows.values())
        c = ws.cell(grand_row, cc, f"=SUM({f.replace('+', ',')})")
        c.number_format = NUM
    for cc in range(1, 2 + len(BLOCKS) * width):
        cel = ws.cell(grand_row, cc)
        cel.fill, cel.font, cel.border = GRAND_FILL, Font(bold=True, size=11), BORDER

    # ブロックの区切りを太線に
    for bi in range(len(BLOCKS)):
        c0 = 2 + bi * width
        for r in range(HR1, grand_row + 1):
            cel = ws.cell(r, c0)
            cel.border = Border(left=MED, right=cel.border.right,
                                top=cel.border.top, bottom=cel.border.bottom)

    ws.column_dimensions["A"].width = 14
    for cc in range(2, 2 + len(BLOCKS) * width):
        ws.column_dimensions[get_column_letter(cc)].width = 14
    ws.row_dimensions[HR2].height = 32
    ws.freeze_panes = ws.cell(FIRST_ROW, 2).coordinate

    note = grand_row + 2
    for i, t in enumerate([
        "※ 振込額＝請求額（税込み）－源泉徴収額。各事務所シートのF列を参照しています。",
        "※ 空欄はその月の金額が0（未計上）であることを表します。",
        "※ 事務所を増やすときは、同じ列構成のシートを追加してこのシートに列を足してください。",
    ]):
        ws.cell(note + i, 1, t).font = Font(size=9, color="808080")

    wb.active = 0
    wb.save(DST)
    print("月数:", len(months), "年計行:", year_rows, "総合計行:", grand_row)


main()
