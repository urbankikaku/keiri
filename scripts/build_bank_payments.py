# -*- coding: utf-8 -*-
"""銀行支払明細ブックを組み直す。

・各月シート  : 1行＝1支払先、横に勘定科目ごとの内訳（202606〜202705の12シート）
・科目別集計  : 何費にいくら払ったかを月次で一覧（全体シート）
・支払先別集計: 支払先ごとの月次推移
・支払先マスタ: 支払先・備考・媒体。ここに追記すると全月シートに行が増える
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC, DST = sys.argv[1], sys.argv[2]

# 元シートの月列（C〜F）→ 月度コード
SRC_COLS = {"202606": 6, "202607": 5, "202608": 4, "202609": 3}

MONTHS = ["202606", "202607", "202608", "202609", "202610", "202611",
          "202612", "202701", "202702", "202703", "202704", "202705"]
MONTH_LABEL = {m: f"{int(m[:4])}年{int(m[4:]):d}月" for m in MONTHS}

ACCOUNTS = ["通信費", "水道光熱費", "賃借料", "リース料", "修繕費", "消耗品費",
            "消耗品費（8%）", "車輛燃料費", "旅費交通費", "接待交際費", "会議費",
            "支払手数料", "租税公課", "保険料", "福利厚生費", "広告宣伝費", "雑費",
            "保険積立金", "工具、器具及び備品", "車両運搬具"]

SPARE_ACCOUNTS = 3            # 科目の追加枠（列）
SPARE_PAYEES = 15             # 支払先の追加枠（行）

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
GRAY = Font(size=9, color="808080")
NUM = "#,##0_);[Red]\\(#,##0\\)"

MASTER = "支払先マスタ"
ACC_SHEET = "科目別集計"
PAY_SHEET = "支払先別集計"


def norm(name):
    """全角/半角カッコと空白の揺れを吸収した突き合わせキー。"""
    t = name.replace("（", "(").replace("）", ")")
    return t.replace("　", "").replace(" ", "")


def collect(ws, bank, rows):
    """親行（勘定科目が空＝小計行）を除いた明細行。支払先・媒体・備考は親から引き継ぐ。"""
    out = []
    parent = {"payee": "", "media": "", "note": ""}
    for r in rows:
        payee = str(ws.cell(r, 1).value or "").strip()
        account = str(ws.cell(r, 2).value or "").strip()
        note = str(ws.cell(r, 7).value or "").strip()
        media = str(ws.cell(r, 8).value or "").strip()
        if payee and not account:
            parent = {"payee": payee, "media": media, "note": note}
            continue
        if not account:
            continue
        if payee:
            parent = {"payee": payee, "media": media, "note": note}
        else:
            payee, media, note = parent["payee"], media or parent["media"], note or parent["note"]
        amounts = {m: (ws.cell(r, c).value if isinstance(ws.cell(r, c).value, (int, float)) else None)
                   for m, c in SRC_COLS.items()}
        out.append({"bank": bank, "payee": payee, "account": account,
                    "note": note, "media": media, "amounts": amounts, "row": r})
    return out


def style_header(cell, wrap=False):
    cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def main():
    wb = openpyxl.load_workbook(SRC)
    items = collect(wb["みずほ銀行"], "みずほ銀行", range(4, 27))
    # 武蔵野銀行 74〜80行は 61〜67行と同一内容の重複のため除外
    items += collect(wb["武蔵野銀行"], "武蔵野銀行", range(2, 74))

    for it in items:
        if it["account"] not in ACCOUNTS:
            ACCOUNTS.append(it["account"])

    # 支払先を（カッコ・空白の揺れを吸収して）一意化。並びは元ブックの登場順
    payees, info = [], {}
    for it in items:
        k = norm(it["payee"])
        if k not in info:
            payees.append(k)
            info[k] = {"name": it["payee"], "banks": [], "paid": [],
                       "note": it["note"], "media": it["media"]}
        d = info[k]
        if it["bank"] not in d["banks"]:
            d["banks"].append(it["bank"])
        if any(it["amounts"].values()) and it["bank"] not in d["paid"]:
            d["paid"].append(it["bank"])
        d["note"] = d["note"] or it["note"]
        d["media"] = d["media"] or it["media"]

    amount = {}
    for it in items:
        for m, v in it["amounts"].items():
            if v:
                key = (norm(it["payee"]), it["account"], m)
                amount[key] = amount.get(key, 0) + v

    n_acc = len(ACCOUNTS) + SPARE_ACCOUNTS
    n_pay = len(payees) + SPARE_PAYEES
    HDR_ROW, FIRST_ROW = 2, 3
    LAST_ROW = FIRST_ROW + n_pay - 1
    TOTAL_ROW = LAST_ROW + 1
    FIRST_COL = 2                      # B列＝最初の勘定科目
    LAST_COL = FIRST_COL + n_acc - 1
    SUM_COL = LAST_COL + 1             # 合計
    NOTE_COL = SUM_COL + 1             # 備考

    for name, new in (("みずほ銀行", "原本_みずほ銀行"), ("武蔵野銀行", "原本_武蔵野銀行")):
        wb[name].title = new
    del wb["勘定科目"]

    # ---------------- 支払先マスタ ----------------
    mst = wb.create_sheet(MASTER)
    mst.cell(1, 1, "支払先マスタ（ここに追記すると全ての月シートに行が増えます）").font = Font(bold=True, size=12)
    for i, h in enumerate(["支払先", "銀行（参考）", "備考", "媒体"], 1):
        style_header(mst.cell(2, i, h))
    for i, k in enumerate(payees):
        r = 3 + i
        d = info[k]
        mst.cell(r, 1, d["name"])
        mst.cell(r, 2, "／".join(d["paid"] or d["banks"]))
        mst.cell(r, 3, d["note"])
        mst.cell(r, 4, d["media"])
    for r in range(3, 3 + n_pay):
        for c in range(1, 5):
            mst.cell(r, c).border = BORDER
    for col, w in [("A", 28), ("B", 22), ("C", 32), ("D", 18)]:
        mst.column_dimensions[col].width = w
    mst.freeze_panes = "A3"

    # ---------------- 月別シート（12枚） ----------------
    for m in MONTHS:
        ws = wb.create_sheet(m)
        t = ws.cell(1, 1, f"{m}（{MONTH_LABEL[m]}）  支払先別 × 勘定科目")
        t.font = Font(bold=True, size=14)
        ws.cell(1, FIRST_COL + 2, "※ 支払先の追加は「支払先マスタ」／科目の追加は「科目別集計」シートで行うと全月に反映されます").font = GRAY

        style_header(ws.cell(HDR_ROW, 1, "支払先"))
        for j in range(n_acc):
            c = ws.cell(HDR_ROW, FIRST_COL + j, f"=IF({ACC_SHEET}!$A{3 + j}=\"\",\"\",{ACC_SHEET}!$A{3 + j})")
            style_header(c, wrap=True)
        style_header(ws.cell(HDR_ROW, SUM_COL, "合計"))
        style_header(ws.cell(HDR_ROW, NOTE_COL, "備考"))
        ws.row_dimensions[HDR_ROW].height = 54

        for i in range(n_pay):
            r = FIRST_ROW + i
            ws.cell(r, 1, f'=IF({MASTER}!A{r}="","",{MASTER}!A{r})').border = BORDER
            for j in range(n_acc):
                acc = ACCOUNTS[j] if j < len(ACCOUNTS) else None
                v = amount.get((payees[i], acc, m)) if (i < len(payees) and acc) else None
                c = ws.cell(r, FIRST_COL + j, v)
                c.number_format, c.border = NUM, BORDER
            L1, L2 = get_column_letter(FIRST_COL), get_column_letter(LAST_COL)
            c = ws.cell(r, SUM_COL, f'=IF(COUNT({L1}{r}:{L2}{r})=0,"",SUM({L1}{r}:{L2}{r}))')
            c.number_format, c.border, c.font = NUM, BORDER, Font(bold=True)
            ws.cell(r, NOTE_COL).border = BORDER

        ws.cell(TOTAL_ROW, 1, "合計")
        for col in range(FIRST_COL, SUM_COL + 1):
            L = get_column_letter(col)
            rng = f"{L}{FIRST_ROW}:{L}{LAST_ROW}"
            c = ws.cell(TOTAL_ROW, col, f'=IF(COUNT({rng})=0,"",SUM({rng}))')
            c.number_format = NUM
        for col in range(1, NOTE_COL + 1):
            c = ws.cell(TOTAL_ROW, col)
            c.fill, c.font, c.border = TOTAL_FILL, Font(bold=True), BORDER

        ws.column_dimensions["A"].width = 28
        for j in range(n_acc):
            ws.column_dimensions[get_column_letter(FIRST_COL + j)].width = 11.5
        ws.column_dimensions[get_column_letter(SUM_COL)].width = 13
        ws.column_dimensions[get_column_letter(NOTE_COL)].width = 26
        ws.freeze_panes = ws.cell(FIRST_ROW, FIRST_COL).coordinate

    # ---------------- 科目別集計（全体） ----------------
    sa = wb.create_sheet(ACC_SHEET, 0)
    sa.cell(1, 1, "勘定科目別 月次支払合計（202606〜202705）").font = Font(bold=True, size=14)
    style_header(sa.cell(2, 1, "勘定科目"))
    for j, m in enumerate(MONTHS):
        style_header(sa.cell(2, 2 + j, m))
    style_header(sa.cell(2, 2 + len(MONTHS), "年間合計"))
    for i in range(n_acc):
        r = 3 + i
        if i < len(ACCOUNTS):
            sa.cell(r, 1, ACCOUNTS[i])
        sa.cell(r, 1).border = BORDER
        col = get_column_letter(FIRST_COL + i)
        for j, m in enumerate(MONTHS):
            c = sa.cell(r, 2 + j, f"='{m}'!{col}{TOTAL_ROW}")
            c.number_format, c.border = NUM, BORDER
        L2 = get_column_letter(1 + len(MONTHS))
        c = sa.cell(r, 2 + len(MONTHS), f"=SUM(B{r}:{L2}{r})")
        c.number_format, c.border, c.font = NUM, BORDER, Font(bold=True)
    tr = 3 + n_acc
    sa.cell(tr, 1, "合計")
    for j in range(len(MONTHS) + 1):
        L = get_column_letter(2 + j)
        c = sa.cell(tr, 2 + j, f"=SUM({L}3:{L}{tr - 1})")
        c.number_format = NUM
    for col in range(1, 3 + len(MONTHS)):
        c = sa.cell(tr, col)
        c.fill, c.font, c.border = TOTAL_FILL, Font(bold=True), BORDER
    for i, t in enumerate([
        "※ 各月シートの合計行から自動集計しています。金額は各月シートに入力してください。",
        "※ A列に科目名を追記すると、全ての月シートの列見出しにも反映されます（追加枠3列）。",
    ]):
        sa.cell(tr + 2 + i, 1, t).font = GRAY
    sa.column_dimensions["A"].width = 24
    for j in range(len(MONTHS) + 1):
        sa.column_dimensions[get_column_letter(2 + j)].width = 12
    sa.freeze_panes = "B3"

    # ---------------- 支払先別集計 ----------------
    sp = wb.create_sheet(PAY_SHEET, 1)
    sp.cell(1, 1, "支払先別 月次支払合計（202606〜202705）").font = Font(bold=True, size=14)
    style_header(sp.cell(2, 1, "支払先"))
    for j, m in enumerate(MONTHS):
        style_header(sp.cell(2, 2 + j, m))
    style_header(sp.cell(2, 2 + len(MONTHS), "年間合計"))
    sum_col_L = get_column_letter(SUM_COL)
    for i in range(n_pay):
        r = 3 + i
        src_row = FIRST_ROW + i
        c = sp.cell(r, 1, f'=IF({MASTER}!A{src_row}="","",{MASTER}!A{src_row})')
        c.border = BORDER
        for j, m in enumerate(MONTHS):
            c = sp.cell(r, 2 + j, f"='{m}'!{sum_col_L}{src_row}")
            c.number_format, c.border = NUM, BORDER
        L2 = get_column_letter(1 + len(MONTHS))
        c = sp.cell(r, 2 + len(MONTHS), f"=SUM(B{r}:{L2}{r})")
        c.number_format, c.border, c.font = NUM, BORDER, Font(bold=True)
    tr2 = 3 + n_pay
    sp.cell(tr2, 1, "合計")
    for j in range(len(MONTHS) + 1):
        L = get_column_letter(2 + j)
        c = sp.cell(tr2, 2 + j, f"=SUM({L}3:{L}{tr2 - 1})")
        c.number_format = NUM
    for col in range(1, 3 + len(MONTHS)):
        c = sp.cell(tr2, col)
        c.fill, c.font, c.border = TOTAL_FILL, Font(bold=True), BORDER
    sp.cell(tr2 + 2, 1, "※ 各月シートの支払先ごとの合計を並べたものです。支払先の追加は「支払先マスタ」で。").font = GRAY
    sp.column_dimensions["A"].width = 28
    for j in range(len(MONTHS) + 1):
        sp.column_dimensions[get_column_letter(2 + j)].width = 12
    sp.freeze_panes = "B3"

    order = [ACC_SHEET, PAY_SHEET] + MONTHS + [MASTER, "原本_みずほ銀行", "原本_武蔵野銀行"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(DST)
    print("支払先:", len(payees), "科目:", len(ACCOUNTS),
          "月シート行:", FIRST_ROW, "-", LAST_ROW, "合計行:", TOTAL_ROW,
          "科目列:", get_column_letter(FIRST_COL), "-", get_column_letter(LAST_COL))


main()
